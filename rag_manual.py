"""Soluciona el RAG de abastecimiento con el método manual de Álvaro.

Uso:  python rag_manual.py RAG.xlsx salida.xlsx [dias_venta] [min_repeticiones]

Hojas 1-6 (se insertan columnas, todo con fórmulas de Excel):
  CONS DIA    = CONSUMO_ACUM_30D / dias_transcurridos   (13 el 14-oct; = fecha BI - 1 día)
  MINIMO      = EXHIBICION
  MAXIMO      = ROUND(MINIMO + EMPAQUE/2, 0)
  DG MAX = MAXIMO/CONS DIA | DIF = MAX-MIN | % Empaque = DIF/EMPAQUE
  SE CONFIGURA= "CORREGIR" si ROQ < EMPAQUE/2, si no "OK"
Hoja 7 (TDF con LTC, sin columnas nuevas): se priorizan los estadísticos que más se repiten
  (>= min_repeticiones casos, por defecto 5; ver hoja DINAMICA). A todas las filas de esos
  productos se les escribe MINIMO=EXHIBICION y MAXIMO=ROUND(MINIMO+EMPAQUE/2). El resto queda
  sin configurar. Se añade la hoja DINAMICA con el conteo de casos por producto.
"""
import sys, pandas as pd, openpyxl
from openpyxl.styles import PatternFill, Font
from copy import copy

AMARILLO = PatternFill('solid', fgColor='FFFF00')

def main(rag, out, dias=None, minrep=5):
    wb = openpyxl.load_workbook(rag)
    resumen = {}
    if dias is None:   # días de venta completos = día de la FECHA - 1
        f = pd.to_datetime(wb['TAREA_1'].cell(4, 1).value, dayfirst=True)
        dias = f.day - 1
    cnt = {}
    w7 = wb['TAREA_7']; h7 = {w7.cell(3, c).value: c for c in range(1, w7.max_column + 1)}
    for r in range(4, w7.max_row + 1):
        k = w7.cell(r, h7['COD_PRODUCTO']).value
        if k is not None:
            cnt[k] = cnt.get(k, 0) + 1
    for ws in wb.worksheets:
        if not ws.title.startswith('TAREA_'):
            continue
        n = int(ws.title.split('_')[1])
        if n < 7:
            ws.insert_cols(16)       # CONS DIA
            ws.insert_cols(19, 4)    # DG MAX, DIF, % Empaque, SE CONFIGURA
            for c, t in zip((16, 19, 20, 21, 22), ('CONS DIA', 'DG MAX', 'DIF', '% Empaque', 'SE CONFIGURA')):
                ws.cell(3, c, t)
                src = ws.cell(3, 15); d = ws.cell(3, c)
                d.font, d.fill, d.border, d.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
            for r in range(4, ws.max_row + 1):
                if ws.cell(r, 6).value is None:
                    continue
                ws[f'P{r}'] = f'=+O{r}/{dias}'
                ws[f'Q{r}'] = f'=+AA{r}'
                ws[f'R{r}'] = f'=+ROUND(Q{r}+(L{r}/2),0)'
                ws[f'S{r}'] = f'=IFERROR(R{r}/P{r},0)'
                ws[f'T{r}'] = f'=+R{r}-Q{r}'
                ws[f'U{r}'] = f'=+T{r}/L{r}'
                ws[f'V{r}'] = f'=+IF(W{r}<(L{r}/2),"CORREGIR","OK")'
            resumen[n] = ws.max_row - 3
        else:
            h = {ws.cell(3, c).value: c for c in range(1, ws.max_column + 1)}
            L = lambda k: openpyxl.utils.get_column_letter(h[k])
            hecho = 0
            for r in range(4, ws.max_row + 1):
                cod = ws.cell(r, h['COD_PRODUCTO']).value
                if cod is None:
                    continue
                if cnt[cod] < minrep:
                    continue
                ws[f'{L("MINIMO")}{r}'] = f'=+{L("EXHIBICION")}{r}'
                ws[f'{L("MAXIMO")}{r}'] = f'=+ROUND({L("MINIMO")}{r}+({L("EMPAQUE")}{r}/2),0)'
                for k in ('MINIMO', 'MAXIMO'):
                    ws[f'{L(k)}{r}'].fill = AMARILLO
                hecho += 1
            resumen[n] = (hecho, ws.max_row - 3)
    # hoja dinámica: casos por producto en TAREA_7
    dn = wb.create_sheet('DINAMICA DE LA HOJA 7')
    dn.append(['COD_PRODUCTO', 'Cantidad de casos'])
    for k, v in sorted(cnt.items(), key=lambda x: -x[1]):
        dn.append([k, v])
    dn.append(['Total', sum(cnt.values())])
    for c in dn[1]:
        c.font = Font(bold=True)
    wb.save(out)
    print('días de venta:', dias, resumen)

if __name__ == '__main__':
    a = sys.argv
    main(a[1], a[2], int(a[3]) if len(a) > 3 else None, int(a[4]) if len(a) > 4 else 5)
